from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from invstruct.export.schema_version import EXPORT_SCHEMA_VERSION
from invstruct.schemas import InvoiceRecord

REPORT_COLUMNS = [
    "file_name",
    "expense_batch_id",
    "doc_id",
    "doc_type",
    "template_id",
    "merchant_name",
    "issue_date",
    "currency",
    "total_amount_gross",
    "tax_amount",
    "invoice_number",
    "tax_number",
    "merchant_tax_id",
    "anomaly_count",
    "anomaly_codes",
    "confidence_overall",
    "status",
    "error_code",
    "error_message",
    "warnings_count",
    "parser_version",
    "retry_key",
    "source_file",
    "sha256",
    "trace_id",
]


def _row(record: InvoiceRecord) -> dict[str, object]:
    anomaly_codes = ",".join(item.code for item in (record.anomalies or []))
    warnings = record.warnings or []
    return {
        "file_name": record.source.file_name,
        "expense_batch_id": None,
        "doc_id": record.doc_id,
        "doc_type": record.doc_type,
        "template_id": record.template_id,
        "merchant_name": record.merchant_name,
        "issue_date": record.issue_date.isoformat() if record.issue_date else None,
        "currency": record.currency,
        "total_amount_gross": record.total_amount_gross,
        "tax_amount": record.tax_amount,
        "invoice_number": record.invoice_number,
        "tax_number": record.merchant_tax_id,
        "merchant_tax_id": record.merchant_tax_id,
        "anomaly_count": len(record.anomalies or []),
        "anomaly_codes": anomaly_codes,
        "confidence_overall": record.confidence.overall,
        "status": record.status,
        "error_code": record.error_code,
        "error_message": record.error_message,
        "warnings_count": len(warnings),
        "parser_version": record.parser_version,
        "retry_key": record.retry_key,
        "source_file": record.source.file_name,
        "sha256": record.source.sha256,
        "trace_id": record.source.trace_id,
    }


def write_report_xlsx(records: list[InvoiceRecord], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "report"
    sheet.append(REPORT_COLUMNS)

    for record in records:
        row = _row(record)
        sheet.append([row[col] for col in REPORT_COLUMNS])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(REPORT_COLUMNS))}{max(sheet.max_row, 1)}"

    for idx, column in enumerate(REPORT_COLUMNS, start=1):
        max_len = len(column)
        for row_idx in range(2, sheet.max_row + 1):
            value = sheet.cell(row=row_idx, column=idx).value
            max_len = max(max_len, len(str(value or "")))
        sheet.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 60)

    meta_sheet = workbook.create_sheet("_meta")
    generated_at = datetime.now(timezone.utc).isoformat()
    meta_sheet["A1"] = "invstruct_export_schema_version"
    meta_sheet["B1"] = EXPORT_SCHEMA_VERSION
    meta_sheet["A2"] = "generated_at"
    meta_sheet["B2"] = generated_at
    meta_sheet["A3"] = "tool_version"
    meta_sheet["B3"] = "invstruct/0.1.0"

    workbook.save(path)
