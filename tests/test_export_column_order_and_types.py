from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from typer.testing import CliRunner

from invstruct.cli import app


def test_export_column_order_and_types(tmp_path: Path) -> None:
    records_path = tmp_path / "records.jsonl"
    old_style = {
        "doc_id": "doc-old",
        "merchant_name": "Store Old",
        "issue_date": "bad-date-string",
        "total_amount_gross": "not-a-number",
        "source": {"file_name": "old.png", "sha256": "a" * 64, "trace_id": "t-old"},
    }
    mixed = {
        "doc_id": "doc-new",
        "merchant_name": "Store New",
        "issue_date": "2026-02-10",
        "total_amount_gross": 88.2,
        "currency": "CNY",
        "invoice_number": "INV-88",
        "merchant_tax_id": "91350100MA12345678",
        "status": "success",
        "error_code": None,
        "error_message": None,
        "warnings": [{"reason": "demo"}],
        "parser_version": "ocr:MockOcrEngine",
        "retry_key": "rk-1",
        "source": {"file_name": "new.png", "sha256": "b" * 64, "trace_id": "t-new"},
    }
    records_path.write_text(
        json.dumps(old_style, ensure_ascii=False) + "\n" + json.dumps(mixed, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    out_csv = tmp_path / "out.csv"
    out_xlsx = tmp_path / "out.xlsx"

    runner = CliRunner()
    result = runner.invoke(app, ["export", str(records_path), "--csv", str(out_csv), "--xlsx", str(out_xlsx)])
    assert result.exit_code == 0, result.stdout

    frame = pd.read_csv(out_csv, comment="#")
    expected_prefix = [
        "file_name",
        "merchant",
        "issue_date",
        "total_amount_gross",
        "currency",
        "invoice_number",
        "tax_number",
    ]
    assert list(frame.columns[: len(expected_prefix)]) == expected_prefix
    assert "status" in frame.columns
    assert "error_code" in frame.columns
    assert "error_message" in frame.columns
    assert "warnings_count" in frame.columns
    assert "parser_version" in frame.columns
    assert "retry_key" in frame.columns

    old_row = frame[frame["file_name"] == "old.png"].iloc[0]
    assert pd.isna(old_row["total_amount_gross"])
    assert str(old_row["status"]).strip() in {"", "success"}

    new_row = frame[frame["file_name"] == "new.png"].iloc[0]
    assert float(new_row["total_amount_gross"]) == 88.2
    assert str(new_row["issue_date"]) == "2026-02-10"

    workbook = load_workbook(out_xlsx)
    header = [workbook.active.cell(1, col).value for col in range(1, workbook.active.max_column + 1)]
    assert "file_name" in header
    assert "error_message" in header
