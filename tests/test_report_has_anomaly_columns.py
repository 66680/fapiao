from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from invstruct.cli import app


def test_report_has_anomaly_columns(tmp_path: Path) -> None:
    records_path = tmp_path / "records.jsonl"
    row_1 = {
        "doc_id": "doc-1",
        "merchant_name": "Store A",
        "issue_date": "2026-02-09",
        "total_amount_gross": 10.0,
        "invoice_number": "INV-1",
        "merchant_tax_id": "91350100MA12345678",
        "confidence": {"overall": 0.95, "fields": {}},
        "source": {"file_name": "a.png", "trace_id": "t-1", "sha256": "a" * 64},
    }
    row_2 = {
        "doc_id": "doc-2",
        "merchant_name": "Store A",
        "issue_date": "2026-02-09",
        "total_amount_gross": 10.0,
        "invoice_number": "INV-1",
        "merchant_tax_id": "91350100MA12345678",
        "confidence": {"overall": 0.95, "fields": {}},
        "source": {"file_name": "b.png", "trace_id": "t-2", "sha256": "a" * 64},
    }
    records_path.write_text(
        json.dumps(row_1, ensure_ascii=False) + "\n" + json.dumps(row_2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    out_xlsx = tmp_path / "report.xlsx"
    runner = CliRunner()
    result = runner.invoke(app, ["export", "--in", str(records_path), "--xlsx", str(out_xlsx)])
    assert result.exit_code == 0, result.stdout

    workbook = load_workbook(out_xlsx)
    sheet = workbook.active
    header = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    assert "anomaly_count" in header
    assert "anomaly_codes" in header
