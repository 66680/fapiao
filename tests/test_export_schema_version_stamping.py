from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from invstruct.cli import app
from invstruct.export.schema_version import EXPORT_SCHEMA_VERSION


def test_export_schema_version_stamping(tmp_path: Path) -> None:
    records = tmp_path / "records.jsonl"
    row = {
        "doc_id": "doc-1",
        "merchant_name": "Store A",
        "issue_date": "2026-02-09",
        "total_amount_gross": 9.9,
        "currency": "CNY",
        "source": {"file_name": "a.png", "sha256": "a" * 64, "trace_id": "t-1"},
    }
    records.write_text(json.dumps(row, ensure_ascii=False) + "\n", encoding="utf-8")

    csv_with_comments = tmp_path / "with_comments.csv"
    xlsx_path = tmp_path / "report.xlsx"
    csv_no_comments = tmp_path / "no_comments.csv"

    runner = CliRunner()
    result = runner.invoke(app, ["export", str(records), "--csv", str(csv_with_comments), "--xlsx", str(xlsx_path)])
    assert result.exit_code == 0, result.stdout
    lines = csv_with_comments.read_text(encoding="utf-8").splitlines()
    assert lines[0].startswith("# invstruct_export_schema_version=")
    assert lines[0].endswith(str(EXPORT_SCHEMA_VERSION))
    assert lines[1].startswith("# generated_at=")

    result_no_comments = runner.invoke(
        app,
        ["export", str(records), "--csv", str(csv_no_comments), "--no-header-comments"],
    )
    assert result_no_comments.exit_code == 0, result_no_comments.stdout
    first_line = csv_no_comments.read_text(encoding="utf-8").splitlines()[0]
    assert not first_line.startswith("#")

    workbook = load_workbook(xlsx_path)
    assert "_meta" in workbook.sheetnames
    meta = workbook["_meta"]
    assert meta["A1"].value == "invstruct_export_schema_version"
    assert int(meta["B1"].value) == EXPORT_SCHEMA_VERSION
