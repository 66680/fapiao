from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from invstruct.exporters.report_xlsx import REPORT_COLUMNS, write_report_xlsx
from invstruct.schemas import Confidence, InvoiceRecord, SourceInfo


def test_export_xlsx_has_columns(tmp_path: Path) -> None:
    record = InvoiceRecord(
        doc_id="doc-1",
        doc_type="receipt",
        issue_date=date(2026, 2, 9),
        merchant_name="Store A",
        currency="CNY",
        total_amount_gross=12.5,
        confidence=Confidence(overall=0.92, fields={"merchant_name": 0.9}),
        source=SourceInfo(file_name="a.png", sha256="a" * 64, trace_id="t-1"),
    )
    output = tmp_path / "report.xlsx"
    write_report_xlsx([record], output)

    workbook = load_workbook(output)
    sheet = workbook.active
    header = [sheet.cell(1, idx + 1).value for idx in range(len(REPORT_COLUMNS))]
    assert header == REPORT_COLUMNS
    assert sheet.freeze_panes == "A2"
